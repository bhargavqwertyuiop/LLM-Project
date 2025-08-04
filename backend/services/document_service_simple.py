import logging
import os
import uuid
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Any
import aiofiles
from fastapi import UploadFile
import PyPDF2
from docx import Document
import asyncio
from core.config import settings

logger = logging.getLogger(__name__)

class SimpleDocumentService:
    """Simplified document service for local development without pandas"""
    
    def __init__(self):
        self.upload_dir = Path(settings.UPLOAD_DIR)
        self.upload_dir.mkdir(exist_ok=True)
        
        # Supported file types (simplified)
        self.supported_extensions = {
            '.pdf': self._extract_pdf_text,
            '.docx': self._extract_docx_text,
            '.doc': self._extract_docx_text,
            '.txt': self._extract_txt_text,
        }
    
    async def save_uploaded_file(self, file: UploadFile) -> Tuple[str, str]:
        """Save uploaded file and return file path and unique ID"""
        # Generate unique filename
        file_id = str(uuid.uuid4())
        file_extension = Path(file.filename).suffix.lower()
        filename = f"{file_id}{file_extension}"
        file_path = self.upload_dir / filename
        
        # Save file
        async with aiofiles.open(file_path, 'wb') as f:
            content = await file.read()
            await f.write(content)
        
        logger.info(f"Saved file: {filename}")
        return str(file_path), file_id
    
    async def process_document(self, file_path: str) -> Dict[str, Any]:
        """Process document and extract text content"""
        file_path_obj = Path(file_path)
        
        if not file_path_obj.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        file_extension = file_path_obj.suffix.lower()
        
        if file_extension not in self.supported_extensions:
            raise ValueError(f"Unsupported file type: {file_extension}")
        
        try:
            # Extract text using appropriate method
            extractor = self.supported_extensions[file_extension]
            text_content = await extractor(file_path)
            
            # Get file metadata
            file_stats = file_path_obj.stat()
            
            result = {
                'file_path': file_path,
                'filename': file_path_obj.name,
                'file_size': file_stats.st_size,
                'file_type': file_extension,
                'text_content': text_content,
                'word_count': len(text_content.split()) if text_content else 0,
                'char_count': len(text_content) if text_content else 0,
                'processing_status': 'success'
            }
            
            logger.info(f"Successfully processed document: {file_path_obj.name}")
            return result
            
        except Exception as e:
            logger.error(f"Error processing document {file_path}: {e}")
            return {
                'file_path': file_path,
                'filename': file_path_obj.name,
                'processing_status': 'error',
                'error_message': str(e)
            }
    
    async def _extract_pdf_text(self, file_path: str) -> str:
        """Extract text from PDF file"""
        text_content = []
        
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    text = page.extract_text()
                    if text.strip():
                        text_content.append(text)
            
            return '\n\n'.join(text_content)
            
        except Exception as e:
            logger.error(f"Error extracting PDF text: {e}")
            raise
    
    async def _extract_docx_text(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        try:
            doc = Document(file_path)
            text_content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = []
                    for cell in row.cells:
                        if cell.text.strip():
                            row_text.append(cell.text)
                    if row_text:
                        text_content.append(' | '.join(row_text))
            
            return '\n\n'.join(text_content)
            
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {e}")
            raise
    
    async def _extract_txt_text(self, file_path: str) -> str:
        """Extract text from TXT file"""
        try:
            async with aiofiles.open(file_path, 'r', encoding='utf-8') as file:
                content = await file.read()
                return content
        except UnicodeDecodeError:
            # Try with different encoding
            async with aiofiles.open(file_path, 'r', encoding='latin-1') as file:
                content = await file.read()
                return content
        except Exception as e:
            logger.error(f"Error extracting TXT text: {e}")
            raise
    
    def get_supported_extensions(self) -> List[str]:
        """Get list of supported file extensions"""
        return list(self.supported_extensions.keys())
    
    async def cleanup_file(self, file_path: str) -> bool:
        """Delete uploaded file"""
        try:
            file_path_obj = Path(file_path)
            if file_path_obj.exists() and file_path_obj.parent == self.upload_dir:
                file_path_obj.unlink()
                logger.info(f"Cleaned up file: {file_path}")
                return True
            return False
        except Exception as e:
            logger.error(f"Error cleaning up file {file_path}: {e}")
            return False
    
    def validate_file(self, file: UploadFile) -> Tuple[bool, str]:
        """Validate uploaded file"""
        # Check file size
        if hasattr(file, 'size') and file.size:
            max_size = settings.MAX_FILE_SIZE_MB * 1024 * 1024
            if file.size > max_size:
                return False, f"File size exceeds {settings.MAX_FILE_SIZE_MB}MB limit"
        
        # Check file extension
        file_extension = Path(file.filename).suffix.lower()
        if file_extension not in self.supported_extensions:
            return False, f"Unsupported file type: {file_extension}. Supported: {', '.join(self.supported_extensions.keys())}"
        
        return True, "Valid file"